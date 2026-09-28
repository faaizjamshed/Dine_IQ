"""Create local presentation and MP4 demo artifacts from verified evidence."""
from __future__ import annotations

import json
import subprocess
from pathlib import Path
from textwrap import wrap

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
import imageio_ffmpeg


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "Documentation" / "submission_media"


def read_json(path: str) -> dict:
    return json.loads((ROOT / path).read_text(encoding="utf-8"))


def font(size: int, bold: bool = False):
    names = ["arialbd.ttf" if bold else "arial.ttf", "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"]
    for name in names:
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default()


def add_ppt_slide(prs: Presentation, title: str, bullets: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    box = slide.shapes.add_textbox(Inches(0.7), Inches(1.35), Inches(12.0), Inches(5.7))
    frame = box.text_frame
    frame.word_wrap = True
    frame.clear()
    for index, bullet in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = bullet
        paragraph.level = 0
        paragraph.font.size = Pt(22)


def draw_slide(path: Path, title: str, bullets: list[str]) -> None:
    image = Image.new("RGB", (1600, 900), "#111827")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 1600, 120), fill="#0f766e")
    draw.text((70, 35), "DineIQ Analytics", fill="white", font=font(36, True))
    draw.text((70, 160), title, fill="#f9fafb", font=font(52, True))
    y = 270
    for bullet in bullets:
        lines = wrap(bullet, 78)
        draw.text((95, y), "-", fill="#5eead4", font=font(30, True))
        for line in lines:
            draw.text((135, y), line, fill="#e5e7eb", font=font(30))
            y += 42
        y += 24
    draw.text((70, 845), "Local evidence demo - no public deployment claim", fill="#9ca3af", font=font(22))
    image.save(path)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    forecast = read_json("results/submission/forecast-v1/evaluation.json")
    comparison = read_json("results/submission/aligned-comparison-v1/demand_model_comparison.json")
    quality = read_json("results/submission/quality/v2/quality_cleaning.json")
    portable = read_json("results/submission/portable_verification.json")
    quality_totals = {
        "raw": sum(table["before_count"] for table in quality["tables"].values()),
        "accepted": sum(table["accepted_count"] for table in quality["tables"].values()),
        "quarantined": sum(table["quarantine_count"] for table in quality["tables"].values()),
    }

    slides = [
        ("Business Problem", ["Restaurant operators need menu, demand, customer, waste and promotion intelligence in one workflow.", "This submission supports verified local analytics with transparent limits."]),
        ("Data Generation", ["11 related synthetic tables cover customers, restaurants, menu, orders, promotions, ratings, inventory and wastage.", "Generation and validation evidence is retained in results/generation."]),
        ("Quality Validation", [f"Strict v2 validation: raw {quality_totals['raw']:,}, accepted {quality_totals['accepted']:,}, quarantined {quality_totals['quarantined']:,}.", "Rules cover schema, domain, derived totals, referential integrity and cascades."]),
        ("Spark Processing", ["CSV sources are ingested and transformed into Parquet analytics through Spark.", "The portable bundle exports verified model/data artifacts so app evaluation does not need HDFS."]),
        ("Menu Intelligence", ["Menu profitability, margins and performance classes are implemented.", "PKR formatting is used in the UI; labels are analytical guidance rather than automated decisions."]),
        ("Customers and Baskets", ["RFM/customer segments and directional basket candidates are available.", "Recommendations are candidates and need business-impact validation before production use."]),
        ("Forecasting", [f"Aligned comparison uses {comparison['test_rows']:,} held-out rows and 500 sampled cases.", f"Daily model beats seasonal baseline at item/category/location levels; item 28-day MAE {forecast['levels']['item']['backtest']['28']['model']['mae']:.2f} vs {forecast['levels']['item']['backtest']['28']['seasonal_naive']['mae']:.2f}."]),
        ("Wastage", ["Wastage outputs summarize cost/quantity and provide weekly risk screening.", "Low precision means review signal only, not operational-grade prediction."]),
        ("Pricing and Promotions", ["Pricing/promotion analysis is descriptive and flags promotion traps.", "It is not a causal uplift engine."]),
        ("App and Security", ["Flask/Waitress serves React with authentication, CSRF, roles and SQLite operational storage.", "Production TLS/secrets are documented but no public host is provisioned."]),
        ("Portable Verification", [f"No-network Docker verification passed: {portable['tests_run']} tests, {portable['failures']} failures, HDFS disabled.", f"Fresh bundle has {read_json('results/submission/serving-v3/manifest.json')['files'].__len__()} checksum-listed files."]),
        ("Evaluator Path", ["Run docker compose -f compose.serving.yml up --build -d.", "Use README, INSTALLATION, EVALUATOR_INSTRUCTIONS, PROJECT_REPORT and DATA_ACCESS for review."]),
    ]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    image_paths = []
    for index, (title, bullets) in enumerate(slides, start=1):
        add_ppt_slide(prs, title, bullets)
        image_path = OUT / f"demo_slide_{index:02d}.png"
        draw_slide(image_path, title, bullets)
        image_paths.append(image_path)
    pptx_path = OUT / "DineIQ_Submission_Presentation.pptx"
    prs.save(pptx_path)

    concat = OUT / "demo_concat.txt"
    lines = []
    for image_path in image_paths:
        lines.append(f"file '{image_path.as_posix()}'")
        lines.append("duration 4")
    lines.append(f"file '{image_paths[-1].as_posix()}'")
    concat.write_text("\n".join(lines) + "\n", encoding="utf-8")
    mp4_path = OUT / "DineIQ_Demo.mp4"
    ffmpeg = imageio_ffmpeg.get_ffmpeg_exe()
    subprocess.run([ffmpeg, "-y", "-f", "concat", "-safe", "0", "-i", str(concat), "-vf", "fps=30,format=yuv420p", str(mp4_path)], check=True)
    print(json.dumps({"pptx": str(pptx_path), "mp4": str(mp4_path), "slides": len(slides)}, indent=2))


if __name__ == "__main__":
    main()
